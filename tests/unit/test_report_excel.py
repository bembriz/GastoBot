from io import BytesIO

from openpyxl import load_workbook

from app.external.layout import COLUMNAS
from app.external.models import ExternalExpense
from app.reports.excel import build_workbook


def _row(**overrides):
    values = {
        "expense_date": "2026-07-15",
        "category": "CAT",
        "description": "OP-1-ticket",
        "employee": "RUBEN BENJAMIN VAZQUEZ EMBRIZ",
        "paid_by": "Empresa",
        "activities": "",
        "accounting_date": "",
        "account": "ACC",
        "unit_price": "123.45",
        "quantity": "1",
        "include_taxes": "",
        "tax_amount": "",
        "total": "123.45",
        "estado": "Por reportar",
        "bank": "BBVA",
        "transaction_type": "Transferencia",
        "source": "app",
        "month_label": "Julio-2026",
    }
    values.update(overrides)
    return ExternalExpense(**values)


def test_workbook_has_month_title_and_headers():
    buffer = build_workbook("Julio-2026", [_row()])
    workbook = load_workbook(BytesIO(buffer.getvalue()))

    assert workbook.sheetnames == ["Julio-2026"]
    sheet = workbook["Julio-2026"]
    assert [cell.value for cell in sheet[1]] == COLUMNAS
    assert sheet.max_column == 16


def test_workbook_rows_match_layout():
    buffer = build_workbook("Julio-2026", [_row()])
    sheet = load_workbook(BytesIO(buffer.getvalue()))["Julio-2026"]

    values = [cell.value for cell in sheet[2]]
    assert values[0] == "2026-07-15"
    assert values[2] == "OP-1-ticket"
    assert values[12] == "123.45"
    assert values[14] == "BBVA"


def test_workbook_without_rows_only_has_headers():
    buffer = build_workbook("Agosto-2026", [])
    sheet = load_workbook(BytesIO(buffer.getvalue()))["Agosto-2026"]
    assert sheet.max_row == 1
