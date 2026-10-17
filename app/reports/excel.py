"""Generacion del libro .xlsx por mes con el layout de 16 columnas."""

from collections.abc import Iterable
from io import BytesIO

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from app.external.layout import CAMPOS, COLUMNAS
from app.external.models import ExternalExpense


def build_workbook(month_label: str, rows: Iterable[ExternalExpense]) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = month_label
    sheet.append(COLUMNAS)
    sheet.freeze_panes = "A2"

    for row in rows:
        sheet.append([getattr(row, field, "") for field in CAMPOS])

    for index, header in enumerate(COLUMNAS, start=1):
        width = max(14, min(45, len(header) + 4))
        sheet.column_dimensions[get_column_letter(index)].width = width

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer
